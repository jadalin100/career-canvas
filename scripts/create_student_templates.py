from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "site" / "templates"
OUTPUT.mkdir(parents=True, exist_ok=True)

NAVY = "1C2E4A"
DUSTY = "52677D"
IVORY = "BDC4D4"
CREAM = "D1CFC9"
LINE = "D9D9D9"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table):
    tbl_pr = table._tbl.tblPr
    borders_el = tbl_pr.find(qn("w:tblBorders"))
    if borders_el is None:
        borders_el = OxmlElement("w:tblBorders")
        tbl_pr.append(borders_el)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders_el.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders_el.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:color"), LINE)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_field_table(doc, rows):
    for label, prompt in rows:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(5)
        label_run = paragraph.add_run(f"{label}: ")
        label_run.bold = True
        label_run.font.size = Pt(10)
        label_run.font.color.rgb = RGBColor(15, 26, 43)
        prompt_run = paragraph.add_run(f"[{prompt}]")
        prompt_run.font.size = Pt(10.5)
        prompt_run.font.color.rgb = RGBColor(82, 103, 125)
    doc.add_paragraph()


def add_writing_space(doc, lines=4):
    paragraph = doc.add_paragraph("[Type your response here]")
    paragraph.paragraph_format.space_after = Pt(28)
    paragraph.runs[0].font.color.rgb = RGBColor(82, 103, 125)
    paragraph.runs[0].font.size = Pt(10.5)


def base_document(title, subtitle):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"].font.color.rgb = RGBColor(15, 26, 43)
    styles["Normal"].paragraph_format.space_after = Pt(7)
    for name, size in (("Title", 20), ("Heading 1", 17), ("Heading 2", 13)):
        styles[name].font.name = "Aptos Display"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
        styles[name].paragraph_format.space_before = Pt(10)
        styles[name].paragraph_format.space_after = Pt(7)
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    p = doc.add_paragraph("CAREER CANVAS")
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(18)
    run = p.runs[0]
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(82, 103, 125)

    doc.add_paragraph(title, style="Title")
    intro = doc.add_paragraph(subtitle)
    intro.paragraph_format.space_after = Pt(14)
    intro.runs[0].font.color.rgb = RGBColor(82, 103, 125)
    intro.runs[0].font.size = Pt(10.5)
    add_field_table(doc, [("Student", "Type your name"), ("Business", "Type the partner business"), ("Date", "Type the date")])

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Career Canvas  |  Great Neck South")
    footer.runs[0].font.size = Pt(9)
    footer.runs[0].font.color.rgb = RGBColor(82, 103, 125)
    return doc


def section(doc, heading, guidance, fields=None, lines=4):
    doc.add_heading(heading, level=1)
    p = doc.add_paragraph(guidance)
    p.runs[0].italic = True
    p.runs[0].font.color.rgb = RGBColor(82, 103, 125)
    if fields:
        add_field_table(doc, fields)
    else:
        add_writing_space(doc, lines)


def save_business_article():
    doc = base_document(
        "Business Article and Campaign Guide",
        "Use this guide to write a one-to-four-page article that explains the partner business and turns your evidence into an original advertisement recommendation.",
    )
    section(doc, "Business introduction", "Explain what the business does, who it serves, and why it matters in the community. Write in your own words.", [
        ("Main product or service", "What the business offers"),
        ("Audience", "Who the business serves"),
        ("Community role", "Why this business matters locally"),
    ])
    section(doc, "Research question", "Write the main question your article will answer.", lines=2)
    section(doc, "Evidence and source check", "Summarize the strongest facts from the business introduction and outside research. For every source, check currency, relevance, authority, accuracy, and purpose.", lines=6)
    doc.add_page_break()
    section(doc, "Audience insight", "Explain one need, preference, or behavior you noticed. Connect it to specific evidence.", lines=5)
    section(doc, "Advertisement recommendation", "Turn the insight into one focused campaign direction.", [
        ("Goal", "What the advertisement should help accomplish"),
        ("Audience", "Who should notice or act"),
        ("Desired action", "What you want the audience to do"),
        ("Key message", "The single idea they should remember"),
        ("Format", "Flyer, social post, or short video"),
    ])
    section(doc, "Creative direction", "Describe the tone, color, type, imagery, and layout. Explain why each choice fits the audience and partner.", lines=5)
    section(doc, "Success measure", "Name one practical way the business could tell whether the idea worked.", lines=3)
    section(doc, "Sources", "List each source with its title, organization, publication date, access date, and link. Do not paste wording from a source into your article.", lines=5)
    doc.save(OUTPUT / "business-article-and-campaign-guide.docx")


def save_portfolio(filename, aesthetic, body_font, subtitle):
    doc = base_document(
        f"{aesthetic} Project Portfolio",
        subtitle,
    )
    doc.styles["Normal"].font.name = body_font
    section(doc, "Project story", "Explain the business challenge, the audience, your first idea, and the evidence that changed your thinking.", lines=6)
    section(doc, "Research to concept", "Show how one research finding shaped your advertisement or quiz.", lines=5)
    section(doc, "Featured work one", "Insert or link a visual. Add a caption that names the goal, design choice, and result.", lines=5)
    doc.add_page_break()
    section(doc, "Featured work two", "Insert or link a second visual. Explain how it connects to the same audience and message.", lines=5)
    section(doc, "Featured work three", "Insert or link a third visual. Show the final quiz, advertisement, article, portfolio page, or resume.", lines=5)
    section(doc, "Feedback and revision", "Name the feedback you received and show the change you made because of it.", lines=4)
    section(doc, "Skills and reflection", "Name the business and creative skills this work proves. Explain what you would improve next.", lines=4)
    doc.save(OUTPUT / filename)


def save_ad_template(filename, title, subtitle, sections):
    doc = base_document(title, subtitle)
    for heading, prompt, lines in sections:
        section(doc, heading, prompt, lines=lines)
    doc.save(OUTPUT / filename)


if __name__ == "__main__":
    save_business_article()
    save_ad_template("ad-flyer-template.docx", "Business Flyer Template", "Plan an original one-page flyer. Use this as a structure, not a finished design.", [
        ("Goal and audience", "What should this flyer accomplish, and who needs to notice it?", 4),
        ("Headline and message", "Write one headline and the single idea the audience should remember.", 4),
        ("Visual plan", "Sketch the placement of the main image, headline, details, and call to action. Name the colors and fonts you plan to use.", 8),
        ("Required information", "List every fact, date, location, link, price, or disclaimer that must be correct.", 4),
        ("Final check", "Can someone understand the message in five seconds? Is every fact verified?", 3),
    ])
    save_ad_template("ad-social-post-template.docx", "Social Media Advertisement Template", "Plan a square or vertical social post and the caption that supports it.", [
        ("Goal and audience", "What should someone think, feel, or do after seeing this post?", 4),
        ("Post layout", "Sketch the visual, headline, logo placement, and call to action.", 8),
        ("Caption", "Write a concise caption in the business's voice. Include a clear next step.", 5),
        ("Accessibility text", "Describe the visual for someone who cannot see it.", 3),
        ("Fact and permission check", "List facts, images, names, or claims that need business approval.", 3),
    ])
    save_ad_template("ad-video-storyboard-template.docx", "Short Video Advertisement Storyboard", "Plan a 15-to-30-second vertical video before recording or editing.", [
        ("Goal and audience", "What should the viewer remember or do?", 3),
        ("Opening hook", "What happens in the first two seconds?", 3),
        ("Six-scene storyboard", "For each scene, write the visual, on-screen words, audio, and approximate time.", 12),
        ("Closing action", "What exact next step appears at the end?", 3),
        ("Production check", "Confirm facts, music rights, image permissions, captions, and business approval.", 4),
    ])
    save_portfolio("portfolio-editorial-blue.docx", "Editorial Blue", "Aptos", "Combine your project story with a visual showcase in a structured editorial style.")
    save_portfolio("portfolio-minimal-ivory.docx", "Minimal Ivory", "Georgia", "Combine your project story with a visual showcase using calm typography and generous space.")
    save_portfolio("portfolio-bold-grid.docx", "Bold Grid", "Trebuchet MS", "Combine your project story with a visual showcase built for strong screenshots and captions.")
    print("Created 7 Career Canvas templates")
